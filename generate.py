# -*- coding: utf-8 -*-
"""请假证明一键生成 Skill v1.0（乙方案：每一页末尾都出现落款）

用法:
    python generate.py <输入文件> [--out <输出目录>] [--name <输出文件名>]

输入支持三种载体（含 名单 + 活动信息）：
    .xlsx  Sheet「名单」(列头=模板列头) + Sheet「活动信息」(键值对: template/event_date/periods/
           location/activity/dorm/special_note/sign_date)
    .csv   名单文件（列头=模板列头），同目录下的 活动信息.csv 存放键值对
    .json  单文件含全部字段（见 样例数据/*.json）

流程：读取输入 → 复制模板 → 填正文黄高亮槽位 → 填名单表格 →
      乙方案分段（每页=表头+行+2空行+落款组+分页符）→ 渲染核验 → 容量自适应收敛 → 输出 docx
"""
import argparse
import json
import math
import os
import re
import subprocess
import sys
from copy import deepcopy
from datetime import date, timedelta

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import qn, nsdecls
from docx.shared import Pt
import pymupdf

WORK = os.path.dirname(os.path.abspath(__file__))
def _slot_ranges_sk(full):
    """上课请假：固定前缀 | S1 日期(星期) | S2 节次 | S3 到{地点}出席{活动}，志愿者活动， | 固定后缀"""
    m1 = re.search(r'\d{4}年\d{1,2}月\d{1,2}日\(星期[一二三四五六日]\)', full)
    m3 = re.search(r'到.+?，志愿者活动，', full)
    if not m1 or not m3:
        raise RuntimeError(f'上课模板正文槽位定位失败: {full!r}')
    return [(m1.start(), m1.end()), (m1.end(), m3.start()), (m3.start(), m3.end())]


def _slot_ranges_fs(full):
    """分寝室：固定 兹证明 | S1 宿舍楼 | 固定 下列名单中同学，因参加 | S2 日期（星期） | S3 参加…， | 固定后缀

    楼号既可能是模板里的中文写法（七舍），也可能是生成后的阿拉伯数字写法（12舍），
    两种都要能定位——否则替换完第一个槽位后，第二个槽位就再也找不到锚点了。
    """
    m1 = re.search(r'[一二三四五六七八九十\d]+舍', full)
    m2 = re.search(r'\d{4}年\d{1,2}月\d{1,2}日（星期[一二三四五六日]）', full)
    end3 = full.find('故需请假，特此证明！')
    if not m1 or not m2 or end3 < 0:
        raise RuntimeError(f'分寝室模板正文槽位定位失败: {full!r}')
    return [(m1.start(), m1.end()), (m2.start(), m2.end()), (m2.end(), end3)]


TEMPLATES = {
    '上课请假': dict(
        file='上课请假模板.docx',
        headers=['学院', '姓名', '班级', '学号'],
        required=['event_date', 'periods', 'location', 'activity'],
        slot_ranges=_slot_ranges_sk,
    ),
    '分寝室请假': dict(
        file='分寝室请假模板.docx',
        headers=['姓名', '班级', '学号', '寝室号'],
        # 楼号不再手填，自动从名单的寝室信息里识别
        required=['event_date', 'activity', 'special_note'],
        slot_ranges=_slot_ranges_fs,
    ),
}
BLANK_LINES = 2
MAX_ROUNDS = 15   # 容量改为单调递减，需给足轮次让它收敛
WEEKDAY_CN = ['星期一', '星期二', '星期三', '星期四', '星期五', '星期六', '星期日']

# ---------------- 节次 → 时间 自动换算 ----------------
# 校历作息：第1-2节 8:00-9:35；第3-4节 10:00-11:35；第5-6节 13:30-15:05；第7-8节 15:30-17:05
# 9-10 / 11-12 节为预留，如与实际不符请改这里。
PERIOD_TIME = {
    (1, 2): ('8:00', '9:35'),
    (3, 4): ('10:00', '11:35'),
    (5, 6): ('13:30', '15:05'),
    (7, 8): ('15:30', '17:05'),
    (9, 10): ('18:00', '19:35'),
    (11, 12): ('19:45', '21:20'),
}
PAIR_OF = {}
for _p in PERIOD_TIME:
    PAIR_OF[_p[0]] = _p
    PAIR_OF[_p[1]] = _p
CN_NUM = {1: '一', 2: '二', 3: '三', 4: '四', 5: '五', 6: '六',
          7: '七', 8: '八', 9: '九', 10: '十', 11: '十一', 12: '十二'}


def parse_periods(text):
    """'7.8节' / '7、8节' / '5-6节' / '第7、8节' / '56节' / '7节' -> [(a,b), ...]"""
    s = str(text or '').strip()
    if not s:
        return []
    nums = []
    for tok in re.findall(r'\d+', s):
        n = int(tok)
        if n > 12 and len(tok) > 1:      # '56' 不可能是第56节，按 5、6 节拆
            nums.extend(int(ch) for ch in tok)
        else:
            nums.append(n)
    nums = sorted(set(n for n in nums if 1 <= n <= 12))
    pairs, used = [], set()
    for n in nums:
        if n in used:
            continue
        p = PAIR_OF.get(n)
        if p and p[0] in nums and p[1] in nums:
            pairs.append(p)
            used.update(p)
        else:                             # 只写单节次：按所属节次对整段处理
            pairs.append((n, n))
            used.add(n)
    return pairs


def build_period_text(values):
    """汇总全部同学的节次，去重排序后合成正文用的时间段串"""
    all_pairs, seen, warns = [], set(), []
    for v in values:
        for p in parse_periods(v):
            if p not in seen:
                seen.add(p)
                all_pairs.append(p)
    all_pairs.sort()
    out = []
    for a, b in all_pairs:
        t = PERIOD_TIME.get((a, b)) or PERIOD_TIME.get(PAIR_OF.get(a, (a, b)))
        span = f'{t[0]}-{t[1]}' if t else ''
        if a == b:
            warns.append(f'只写了第{a}节，按 {PAIR_OF.get(a,(a,a))[0]}-{PAIR_OF.get(a,(a,a))[1]} 节整段处理')
            out.append(f'第{CN_NUM.get(a, a)}节课（{span}）')
        else:
            out.append(f'{CN_NUM.get(a, a)}{CN_NUM.get(b, b)}节课（{span}）')
    for w in dict.fromkeys(warns):
        print(f'  [提示] {w}')
    return '、'.join(out)


# ---------------- 寝室信息规范化 ----------------

CN_DIGITS = {'一': 1, '二': 2, '三': 3, '四': 4, '五': 5,
             '六': 6, '七': 7, '八': 8, '九': 9}
BUILDING_RE = re.compile(r'(\d+|[一二三四五六七八九十]+)\s*(?:舍|号楼|号公寓|号宿舍楼|公寓|栋)')
DISTRICT_RE = re.compile(r'[一二三四五六七八九十\d]+\s*区')
FULLWIDTH_DIGITS = str.maketrans('０１２３４５６７８９', '0123456789')


def cn_num_to_int(text):
    """中文数字转整数：'十二'->12、'二十'->20、'八'->8；阿拉伯数字原样返回"""
    t = str(text or '').strip().translate(FULLWIDTH_DIGITS)
    if not t:
        return None
    if t.isdigit():
        return int(t)
    total, section = 0, 0
    for ch in t:
        if ch in CN_DIGITS:
            section = CN_DIGITS[ch]
        elif ch == '十':
            section = (section or 1) * 10
            total += section
            section = 0
        else:
            return None
    return total + section


def parse_dorm(raw):
    """
    解析原始寝室信息，返回 (楼号int|None, 寝室号str|None)。

    各组织表单写法不统一，这里统一规范：
      '8舍'                -> (8, None)        只有楼号，缺寝室号
      '十二舍二区'          -> (12, None)       区号丢弃
      '十二舍一区二区301'    -> (12, '301')     多区号一并丢弃
      '12舍301'            -> (12, '301')
      '8号楼3层301'         -> (8, '301')
      '301'                -> (None, '301')    只有寝室号，缺楼号
    """
    s = str(raw or '').strip().translate(FULLWIDTH_DIGITS)
    if not s:
        return None, None

    building = None
    m = BUILDING_RE.search(s)
    if m:
        building = cn_num_to_int(m.group(1))
        s = s[:m.start()] + ' ' + s[m.end():]

    s = DISTRICT_RE.sub(' ', s)          # 楼区号一律丢弃

    rooms = re.findall(r'\d+', s)
    room = max(rooms, key=len) if rooms else None
    return building, room


def norm_building(n):
    """楼号 -> 统一写法 '12舍'"""
    return f'{n}舍' if n is not None else ''


# ---------------- 输入读取 ----------------

def read_input(path):
    ext = path.rsplit('.', 1)[-1].lower()
    if ext == 'json':
        with open(path, encoding='utf-8') as f:
            data = json.load(f)
        return data.get('info', data), data.get('headers') or data.get('header'), data.get('students', [])
    if ext in ('xlsx', 'xls'):
        return _read_xlsx(path)
    if ext == 'csv':
        return _read_csv(path)
    raise RuntimeError(f'不支持的输入类型: {ext}（支持 xlsx/csv/json）')


def _read_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    names = wb.sheetnames
    sheet_n = next((s for s in names if '名单' in s or '学生' in s), names[0])
    sheet_i = next((s for s in names if '活动' in s or '信息' in s), None)
    rows = [r for r in wb[sheet_n].iter_rows(values_only=True)]
    header = [str(c).strip() for c in rows[0] if c is not None]
    students = [list(r) for r in rows[1:] if any(c is not None and str(c).strip() for c in r)]
    info = {}
    if sheet_i:
        for r in wb[sheet_i].iter_rows(values_only=True):
            if r[0] is None or str(r[0]).strip() == '':
                continue
            k = str(r[0]).strip()
            v = r[1] if len(r) > 1 else None
            if v is not None:
                info[k] = str(v).strip() if not isinstance(v, (int, float)) else str(v)
    return info, header, students


def _read_csv(path):
    import csv
    with open(path, encoding='utf-8-sig') as f:
        rows = [r for r in csv.reader(f)]
    header = [c.strip() for c in rows[0]]
    students = [r for r in rows[1:] if any(c.strip() for c in r)]
    info = {}
    import glob
    d = os.path.dirname(path) or '.'
    base = os.path.basename(path)
    prefix = base.split('-名单')[0] if '-名单' in base else None
    cands = []
    if prefix:
        cands = sorted(glob.glob(os.path.join(d, f'活动信息-{prefix}*.csv')))
    if not cands:
        cands = sorted(glob.glob(os.path.join(d, '活动信息*.csv')))
    seen = set()
    for info_path in cands:
        if info_path in seen:
            continue
        seen.add(info_path)
        with open(info_path, encoding='utf-8-sig') as f:
            for r in csv.reader(f):
                if r and r[0].strip():
                    info[r[0].strip()] = r[1].strip() if len(r) > 1 else ''
    return info, header, students


def parse_date(s):
    s = str(s).strip().replace('年', '-').replace('月', '-').replace('日', '')
    parts = [int(x) for x in re.split(r'[-/.]', s) if x]
    return date(parts[0], parts[1], parts[2])


def fmt_cn_date(d):
    return f'{d.year}年{d.month}月{d.day}日'


# ---------------- 模板编辑 ----------------

def set_flag(p_el, tag):
    pPr = p_el.find(qn('w:pPr'))
    if pPr is None:
        pPr = parse_xml(f'<w:pPr {nsdecls("w")}/>')
        p_el.insert(0, pPr)
    if pPr.find(qn(tag)) is None:
        pPr.append(parse_xml(f'<{tag} {nsdecls("w")}/>'))


def em_width(s):
    """文本的视觉宽度（单位 em）：中日韩全角字符 1.0，半角数字/字母 0.5"""
    w = 0.0
    for ch in str(s or ''):
        w += 1.0 if ord(ch) > 0x2E80 else 0.5
    return w


def layout_table(tbl, headers, rows, section, max_size=16, min_size=12):
    """
    按「最长内容」重算列宽，保证每一列都只占一行，且不删减任何真实信息。
    放不下时优先降字号（16→15→14→13→12pt），而不是截断或换行。
    """
    usable = int((section.page_width - section.left_margin - section.right_margin) / 635)  # twips
    ncol = len(headers)
    cell_mar = 108                     # 单元格左右内边距各 108 twips
    margins = cell_mar * 2

    max_em = []
    for j in range(ncol):
        m = em_width(headers[j])
        for r in rows:
            if j < len(r):
                m = max(m, em_width(r[j]))
        max_em.append(m * 1.02 + 0.4)  # 留一点安全余量，避免字宽估算误差导致换行

    size, need = None, None
    for S in range(max_size, min_size - 1, -1):
        cand = [int(round(e * S * 20)) + margins for e in max_em]
        if sum(cand) <= usable:
            size, need = S, cand
            break
    if size is None:                   # 12pt 仍放不下：等比压缩（极端长文本兜底）
        size = min_size
        need = [int(round(e * size * 20)) + margins for e in max_em]
        scale = usable / sum(need)
        need = [max(int(n * scale), margins + 60) for n in need]
    slack = usable - sum(need)
    if slack > 0:                      # 余量均分，视觉上更舒展
        add = slack // ncol
        need = [n + add for n in need]

    # 固定布局：禁止 Word/WPS 自动调整，否则列宽会被改回去
    tblPr = tbl.tblPr
    for old in tblPr.findall(qn('w:tblLayout')):
        tblPr.remove(old)
    tblPr.append(parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>'))
    for old in tblPr.findall(qn('w:tblW')):
        tblPr.remove(old)
    tblPr.append(parse_xml(f'<w:tblW {nsdecls("w")} w:type="dxa" w:w="{sum(need)}"/>'))

    grid = tbl.find(qn('w:tblGrid'))
    for gc, w in zip(grid.findall(qn('w:gridCol')), need):
        gc.set(qn('w:w'), str(w))

    for tr in tbl.findall(qn('w:tr')):
        for tc, w in zip(tr.findall(qn('w:tc')), need):
            tcPr = tc.find(qn('w:tcPr'))
            if tcPr is None:
                tcPr = parse_xml(f'<w:tcPr {nsdecls("w")}/>')
                tc.insert(0, tcPr)
            for old in tcPr.findall(qn('w:tcW')):
                tcPr.remove(old)
            tcPr.insert(0, parse_xml(f'<w:tcW {nsdecls("w")} w:type="dxa" w:w="{w}"/>'))
            if tcPr.find(qn('w:tcMar')) is None:
                tcPr.append(parse_xml(
                    f'<w:tcMar {nsdecls("w")}>'
                    f'<w:left w:type="dxa" w:w="{cell_mar}"/>'
                    f'<w:right w:type="dxa" w:w="{cell_mar}"/></w:tcMar>'))
    return size


def make_runs(p, values, sz_half=32, red=False):
    for r in p.findall(qn('w:r')):
        p.remove(r)
    for val in values:
        color = '<w:color w:val="FF0000"/>' if red else ''
        r = parse_xml(
            f'<w:r {nsdecls("w")}><w:rPr><w:rFonts w:hint="eastAsia" '
            f'w:ascii="仿宋" w:hAnsi="仿宋" w:eastAsia="仿宋"/>'
            f'{color}'
            f'<w:sz w:val="{sz_half}"/><w:szCs w:val="{sz_half}"/></w:rPr>'
            f'<w:t xml:space="preserve">{val}</w:t></w:r>')
        p.append(r)


def set_date(ps, date_str):
    runs = ps.findall(qn('w:r'))
    if len(runs) >= 2:
        r2 = runs[1]
        ts = r2.findall(qn('w:t'))
        if ts:
            ts[0].text = date_str
            ts[0].set(qn('xml:space'), 'preserve')
        else:
            r2.append(parse_xml(f'<w:t {nsdecls("w")} xml:space="preserve">{date_str}</w:t>'))
        for extra in runs[2:]:
            ps.remove(extra)
    elif runs:
        nr = deepcopy(runs[0])
        ts = nr.findall(qn('w:t'))
        if ts:
            ts[0].text = date_str
            ts[0].set(qn('xml:space'), 'preserve')
        ps.append(nr)


def _set_run_text(r_el, text):
    for t in r_el.findall(qn('w:t')):
        r_el.remove(t)
    r_el.append(parse_xml(f'<w:t {nsdecls("w")} xml:space="preserve">{text}</w:t>'))


def fill_body(body_p, slot_values, slot_ranges):
    """按模板语义槽位（正则定位文本区间）替换正文，保留固定 run 原位置。
    槽位从左到右依次替换；每次替换后重算全文与区间（槽位值长度与原文本
    不同会使后续偏移失效，必须重算），并清理空 run。"""
    p_el = body_p._p
    for slot_idx, new_val in enumerate(slot_values):
        runs = p_el.findall(qn('w:r'))
        texts = [''.join(t.text or '' for t in r.findall(qn('w:t'))) for r in runs]
        offsets, pos = [], 0
        for t in texts:
            offsets.append(pos)
            pos += len(t)
        full = ''.join(texts)
        ranges = slot_ranges(full)
        if len(ranges) != len(slot_values):
            raise RuntimeError(f'槽位数量不匹配: 模板={len(ranges)} 输入={len(slot_values)}')
        s, e = ranges[slot_idx]
        idxs = [i for i in range(len(runs))
                if offsets[i] < e and offsets[i] + len(texts[i]) > s]
        if not idxs:
            raise RuntimeError(f'槽位区间 [{s},{e}) 未命中任何 run')
        i0, i1 = idxs[0], idxs[-1]
        if i0 == i1:
            prefix = texts[i0][: s - offsets[i0]]
            suffix = texts[i0][e - offsets[i0]:]
            _set_run_text(runs[i0], prefix + new_val + suffix)
        else:
            prefix = texts[i0][: s - offsets[i0]]
            suffix = texts[i1][e - offsets[i1]:]
            _set_run_text(runs[i0], prefix + new_val)
            _set_run_text(runs[i1], suffix)
            for i in range(i0 + 1, i1):
                runs[i].getparent().remove(runs[i])
            if suffix == '':
                runs[i1].getparent().remove(runs[i1])
        # 清理空 run
        for r in list(p_el.findall(qn('w:r'))):
            if not ''.join(t.text or '' for t in r.findall(qn('w:t'))):
                r.getparent().remove(r)


def build(src, dst, students, segs, slot_values, sign_date, slot_ranges,
          blank_lines=2, red_cells=()):
    doc = Document(src)
    body = doc.element.body
    src_tbl = doc.tables[0]._tbl

    # 按全部名单的最长内容重算列宽（所有页用同一套宽度），保证每列只占一行
    hdr_texts = [c.text.strip() for c in doc.tables[0].rows[0].cells]
    data_size = layout_table(src_tbl, hdr_texts, students, doc.sections[0])

    # 正文段 = 含「兹证明」的段落
    body_p = next(p for p in doc.paragraphs if '兹证明' in p.text)
    fill_body(body_p, slot_values, slot_ranges)

    paras = doc.paragraphs
    date_re = re.compile(r'\s*\d{4}年\d{1,2}月\d{1,2}日\s*$')
    di = [i for i, p in enumerate(paras) if date_re.match(p.text)][-1]
    date_p = paras[di]
    sig_ps = [paras[di - 2], paras[di - 1], date_p]
    blanks_ps = [p for p in paras[2:di - 2] if p.text.strip() == ''] or [paras[2]]

    src_tbl.getparent().remove(src_tbl)
    for p in blanks_ps + sig_ps:
        p._p.getparent().remove(p._p)

    anchor = paras[1]._p
    gi = 0                     # 学生在整份名单中的序号，用于定位需要标红的单元格
    for si, seg in enumerate(segs):
        new_tbl = deepcopy(src_tbl)
        rows = new_tbl.findall(qn('w:tr'))
        header_tr = rows[0]
        data_trs = rows[1:]
        # 模板预留空行有限（上课 10 / 分寝室 12）。若本段行数超过预留行数，
        # 必须克隆补行，否则超出的学生会被静默丢弃（student_iter 取不到值→数据丢失）。
        while len(data_trs) < seg:
            extra = deepcopy(data_trs[-1])
            new_tbl.append(extra)
            data_trs.append(extra)
        for tr in data_trs[seg:]:
            new_tbl.remove(tr)
        trPr0 = header_tr.find(qn('w:trPr'))
        if trPr0 is not None and trPr0.find(qn('w:tblHeader')) is None:
            trPr0.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
        for tr in data_trs[:seg]:
            tcs = tr.findall(qn('w:tc'))
            vals = students[gi]
            for ci, tc in enumerate(tcs[:len(vals)]):
                make_runs(tc.find(qn('w:p')), [str(vals[ci])],
                          sz_half=data_size * 2, red=(gi, ci) in red_cells)
            gi += 1
        anchor.addnext(new_tbl)
        anchor = new_tbl

        for _ in range(blank_lines):
            pb = deepcopy(blanks_ps[0]._p)
            set_flag(pb, 'w:keepNext')
            anchor.addnext(pb)
            anchor = pb
        for idx, sp in enumerate(sig_ps):
            ps = deepcopy(sp._p)
            set_flag(ps, 'w:keepLines')
            if idx < 2:
                set_flag(ps, 'w:keepNext')
            if sp is date_p:
                set_date(ps, sign_date)
                if si < len(segs) - 1:
                    # 分页符内嵌日期段末尾：不产生独立段落，杜绝空页/分页段孤行
                    ps.append(parse_xml(f'<w:r {nsdecls("w")}><w:br w:type="page"/></w:r>'))
            anchor.addnext(ps)
            anchor = ps

    # 黄色高亮只是模板里「这里要改」的标记，成品必须去掉
    for hl in list(doc.element.body.iter(qn('w:highlight'))):
        hl.getparent().remove(hl)

    doc.save(dst)


# ---------------- 渲染与核验 ----------------

def _set_font(run, name='仿宋', size_pt=12, bold=False):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = parse_xml(f'<w:rFonts {nsdecls("w")}/>')
        rpr.insert(0, rfonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia'):
        rfonts.set(qn(attr), name)


INCOMPLETE_HEADERS = ['姓名', '班级', '学号', '原填写内容', '缺失项']


def write_incomplete_report(path, records, title='信息不全人员名单（缺寝室楼号，需人工补充）'):
    """
    把信息不全的人单独列成一张表，便于二次统计补录。
    records: [{'姓名':..,'班级':..,'学号':..,'原填写内容':..,'缺失项':..}, ...]
    """
    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name = '仿宋'
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), '仿宋')

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(p.add_run(title), '黑体', 16, True)

    tip = doc.add_paragraph()
    tip.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(tip.add_run(f'共 {len(records)} 人，请补齐寝室楼号后重新生成正式证明。'), '仿宋', 12)

    tbl = doc.add_table(rows=1, cols=len(INCOMPLETE_HEADERS))
    tbl.style = 'Table Grid'
    for i, h in enumerate(INCOMPLETE_HEADERS):
        cell = tbl.rows[0].cells[i]
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_font(cell.paragraphs[0].add_run(h), '仿宋', 14, True)
    for rec in records:
        cells = tbl.add_row().cells
        for i, h in enumerate(INCOMPLETE_HEADERS):
            cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            _set_font(cells[i].paragraphs[0].add_run(str(rec.get(h, ''))), '仿宋', 14)
    doc.save(path)
    return path


def render_pdf(docx_path):
    # Word COM 以自身进程目录（常为 C:\WINDOWS\system32）解析相对路径，
    # 必须先把路径转成绝对路径，否则报“找不到文件”。
    docx_path = os.path.abspath(docx_path)
    pdf = docx_path[:-5] + '.pdf'
    # 优先 Word，装了 WPS 的机器回退到 KWps.Application（WPS 兼容 Word COM 接口）
    progid_fallback = (
        "  $word = $null\n"
        "  foreach ($prog in @('Word.Application','KWps.Application','Wps.Application')) {\n"
        "    try { $word = New-Object -ComObject $prog; break } catch {}\n"
        "  }\n"
        "  if ($word -eq $null) { throw '未找到可用的 Word/WPS COM 接口' }\n"
    )
    # 中文 Windows 默认 GBK，统一成 UTF-8 输出，避免子进程解码崩溃
    ps = (
        "[Console]::OutputEncoding=[System.Text.Encoding]::UTF8\n"
        "$ErrorActionPreference = 'Stop'\n"
        "try {\n"
        + progid_fallback
        + "  $word.Visible = $false\n"
        "  $word.DisplayAlerts = 0\n"
        f"  $doc = $word.Documents.Open('{docx_path}', $false, $true)\n"
        f"  $doc.ExportAsFixedFormat('{pdf}', 17)\n"
        "  $doc.Close($false)\n"
        "  $word.Quit()\n"
        "} catch {\n"
        "  Write-Output ('RENDER_ERR: ' + $_.Exception.Message)\n"
        "  try { $word.Quit() } catch {}\n"
        "  exit 1\n"
        "}\n"
    )
    r = subprocess.run(['powershell', '-NoProfile', '-Command', ps],
                       capture_output=True, text=True,
                       encoding='utf-8', errors='replace')
    if r.returncode != 0:
        raise RuntimeError('Word 渲染失败: ' + (r.stdout or '') + (r.stderr or ''))
    return pdf


def verify_pdf(pdf):
    doc = pymupdf.open(pdf)
    problems, infos = [], []
    for i, page in enumerate(doc):
        blocks = [b for b in page.get_text('blocks') if b[6] == 0 and b[4].strip()]
        blocks.sort(key=lambda b: b[3])
        text = ''.join(b[4] for b in blocks).replace(' ', '')
        has_sig = ('共青团' in text and '服务站' in text
                   and re.search(r'\d{4}年\d{1,2}月\d{1,2}日', text) is not None)
        hdr_blocks = [b for b in page.get_text('blocks') if ('学号' in b[4] or '寝室号' in b[4])]
        has_table = len(hdr_blocks) > 0
        hdr_bottom = max(b[3] for b in hdr_blocks) if hdr_blocks else 0
        rows_cnt = len([b for b in page.get_text('blocks')
                        if b[0] < 150 and b[1] > hdr_bottom and b[4].strip()])
        imgs = len(page.get_images(full=True))
        last = blocks[-1][4].strip().replace(' ', '') if blocks else ''
        if not blocks:
            problems.append((i, 'EMPTY', rows_cnt))
        if not has_sig:
            problems.append((i, 'NO_SIG', rows_cnt))
        if has_sig and not has_table:
            problems.append((i, 'ORPHAN', rows_cnt))
        if imgs < 1:
            problems.append((i, 'NO_HEADER', rows_cnt))
        if not (last and '年' in last and '日' in last):
            problems.append((i, 'NO_DATE_END', rows_cnt))
        infos.append(dict(page=i + 1, has_sig=has_sig, has_table=has_table,
                          rows=rows_cnt, last_y=round(blocks[-1][3], 1) if blocks else 0))
    doc.close()
    return problems, infos


def compute_segs(n, F1, F2):
    if n <= 0:
        raise ValueError('名单为空')
    if n <= F1:
        return [n]
    segs = [F1]
    rest = n - F1
    npg = math.ceil(rest / F2)
    base, rem = divmod(rest, npg)
    segs += [base + (1 if i < rem else 0) for i in range(npg)]
    return segs


def solve(src, dst, students, slot_values, sign_date, slot_ranges,
          init_caps=(7, 11), red_cells=()):
    F1, F2 = init_caps
    last_pdf = None
    for rnd in range(1, MAX_ROUNDS + 1):
        segs = compute_segs(len(students), F1, F2)
        build(src, dst, students, segs, slot_values, sign_date, slot_ranges,
              BLANK_LINES, red_cells)
        pdf = render_pdf(dst)
        last_pdf = pdf
        problems, infos = verify_pdf(pdf)
        log = [f'[R{rnd}] caps=({F1},{F2}) segs={segs} pages={len(infos)}']
        for inf in infos:
            log.append(f'  P{inf["page"]}: rows={inf["rows"]:2d} last_y={inf["last_y"]:6.1f}')
        if not problems:
            log.append('  核验全部通过 ✓')
            print('\n'.join(log))
            return segs
        # 注意：verify_pdf 统计的 rows 含正文折行等左对齐文本块，不是真实表格行数，
        # 用它反推容量会得到偏大的值，配合 min() 会导致容量永不下降、6 轮死循环。
        # 改为「每轮单调递减 1」：只要首页有问题就 F1-1，其余页有问题就 F2-1。
        # 必须每轮只减一次（而不是每个问题页各减一次），否则多页同时报错时会
        # 一轮内暴跌，收敛到「每页 1 人」这种退化结果。
        dec_f1 = dec_f2 = False
        for i, kind, rows in problems:
            if kind in ('NO_SIG', 'NO_DATE_END'):
                if i == 0:
                    dec_f1 = True
                else:
                    dec_f2 = True
                log.append(f'  BAD P{i + 1}: {kind} (rows={rows})')
            elif kind in ('ORPHAN', 'EMPTY'):
                if (i - 1) == 0:
                    dec_f1 = True
                else:
                    dec_f2 = True
                log.append(f'  BAD P{i + 1}: {kind} (rows={rows})')
            else:
                log.append(f'  WARN P{i + 1}: {kind}')
        if dec_f1:
            F1, old = max(F1 - 1, 1), F1
            log.append(f'  ADJ F1 {old}->{F1}')
        if dec_f2:
            F2, old = max(F2 - 1, 1), F2
            log.append(f'  ADJ F2 {old}->{F2}')
        print('\n'.join(log))
    raise RuntimeError('多次调整仍无法通过分页核验，请检查模板或数据')


# ---------------- 主流程 ----------------

# 各组织收集的表单列名不统一，按别名定位
COLUMN_ALIASES = {
    '姓名': ['姓名', '名字', '学生姓名', '同学'],
    '班级': ['班级', '行政班级', '所在班级', '班'],
    '学号': ['学号', '学生证号', '学籍号'],
    '学院': ['学院', '所在学院', '院系', '二级学院'],
    '寝室号': ['寝室号', '寝室', '宿舍', '宿舍号', '寝室信息', '宿舍信息', '房间号', '住所'],
    '节次': ['节次', '请假节次', '请假课时', '上课节次', '请假时间', '课程节次'],
}


def resolve_column(header, field):
    """按别名把模板字段映射到源表列下标，精确优先、包含次之"""
    aliases = COLUMN_ALIASES.get(field, [field])
    for al in aliases:
        for i, h in enumerate(header):
            if h == al:
                return i
    for al in aliases:
        for i, h in enumerate(header):
            if h and al in h:
                return i
    return None


def find_template(tpl):
    for cand in (os.path.join(WORK, '模板', tpl['file']), os.path.join(WORK, tpl['file'])):
        if os.path.exists(cand):
            return cand
    raise RuntimeError(f'模板文件不存在: {tpl["file"]}（需随包放在 模板/ 子目录或与 generate.py 同目录）')


def main():
    ap = argparse.ArgumentParser(description='请假证明一键生成（乙方案）')
    ap.add_argument('input', help='输入文件 xlsx/csv/json')
    ap.add_argument('--out', default=WORK, help='输出目录（默认脚本目录）')
    ap.add_argument('--name', default=None, help='输出文件名（默认自动）')
    args = ap.parse_args()

    info, header, students_raw = read_input(args.input)

    template_id = info.get('template') or info.get('模板')
    if not template_id:
        if resolve_column(header, '学院') is not None:
            template_id = '上课请假'
        elif resolve_column(header, '寝室号') is not None:
            template_id = '分寝室请假'
    if template_id not in TEMPLATES:
        raise RuntimeError(f'未知模板: {template_id}（可选: {list(TEMPLATES)}）')
    tpl = TEMPLATES[template_id]

    # 列定位（允许别名）
    col_idx = {}
    for h in tpl['headers']:
        i = resolve_column(header, h)
        if i is None:
            raise RuntimeError(f'名单中找不到「{h}」列（表头为 {header}）')
        col_idx[h] = i

    ev = parse_date(info['event_date'])
    sign = parse_date(info.get('sign_date') or info.get('落款日期') or (ev - timedelta(days=1)))
    wd = WEEKDAY_CN[ev.weekday()]
    sign_date = fmt_cn_date(sign)
    src = find_template(tpl)
    stamp = dict(y=ev.year, m=ev.month, d=ev.day, wd=wd)

    if template_id == '上课请假':
        # 「请假节次」列 → 正文时间段（该列不进表格，但是正文驱动列，不能丢）
        pt = resolve_column(header, '节次')
        if not info.get('periods') and pt is not None:
            info['periods'] = build_period_text(
                [r[pt] if pt < len(r) else '' for r in students_raw])
            print(f'[自动换算] 请假节次 -> {info["periods"]}')
        missing = [f for f in tpl['required'] if not info.get(f)]
        if missing:
            raise RuntimeError(f'活动信息缺少字段: {missing}')

        students = [[row[col_idx[h]] if col_idx[h] < len(row) else ''
                     for h in tpl['headers']] for row in students_raw]
        slot_values = [f'{stamp["y"]}年{stamp["m"]}月{stamp["d"]}日({stamp["wd"]})',
                       str(info['periods']).rstrip() + ' ',
                       f"到{info['location']}出席{info['activity']}，志愿者活动，"]
        name = args.name or f'上课请假证明-{ev:%Y%m%d}.docx'
        dst = os.path.join(args.out, name)
        solve(src, dst, students, slot_values, sign_date, tpl['slot_ranges'])
        print(f'\n输出: {dst}')
        print(f'落款日期: {sign_date}  |  学生数: {len(students)}')
        return [dst]

    # ---------------- 分寝室请假：按楼号分组，每栋楼一份 ----------------
    return run_dorm(src, args, info, tpl, col_idx, students_raw, stamp, sign_date)


def run_dorm(src, args, info, tpl, col_idx, students_raw, stamp, sign_date):
    missing = [f for f in tpl['required'] if not info.get(f)]
    if missing:
        raise RuntimeError(f'活动信息缺少字段: {missing}')

    # 寝室列：楼号 + 寝室号 拆开，楼号统一为「数字+舍」，区号丢弃
    di = col_idx['寝室号']

    def cell(row, i):
        v = row[i] if i < len(row) else ''
        return '' if v is None else str(v).strip()

    groups, incomplete, raw_of = {}, [], {}
    for i, row in enumerate(students_raw):
        raw = cell(row, di)
        building, room = parse_dorm(raw)
        raw_of[i] = raw
        if building is None:
            incomplete.append((i, raw, room))
            continue
        groups.setdefault(building, []).append((i, room))

    if not groups and not incomplete:
        raise RuntimeError('名单里没有任何可用的寝室信息')

    outputs = []
    for building in sorted(groups):
        members = groups[building]
        students, red_cells = [], set()
        ci_room = tpl['headers'].index('寝室号')
        for local_i, (gi, room) in enumerate(members):
            row = students_raw[gi]
            vals = [cell(row, col_idx[h]) for h in tpl['headers']]
            if room:
                vals[ci_room] = room
            else:
                # 只有楼号没有寝室号：保留原填写内容并标红，提示人工复检
                vals[ci_room] = raw_of[gi]
                red_cells.add((local_i, ci_room))
            students.append(vals)

        slot_values = [norm_building(building),
                       f'{stamp["y"]}年{stamp["m"]}月{stamp["d"]}日（{stamp["wd"]}）',
                       f"参加{info['activity']}志愿者工作，{info['special_note']}，"]
        default_name = (f'分寝室请假证明-{norm_building(building)}-'
                        f'{stamp["y"]:04d}{stamp["m"]:02d}{stamp["d"]:02d}.docx')
        if args.name:
            # 多个楼号时必须在文件名里区分，否则会互相覆盖
            stem, ext = os.path.splitext(args.name)
            name = f'{stem}-{norm_building(building)}{ext or ".docx"}' if len(groups) > 1 else args.name
        else:
            name = default_name
        dst = os.path.join(args.out, name)
        print(f'\n=== {norm_building(building)}：{len(students)} 人'
              f'{"（含 " + str(len(red_cells)) + " 人缺寝室号，已标红）" if red_cells else ""} ===')
        solve(src, dst, students, slot_values, sign_date, tpl['slot_ranges'], red_cells=red_cells)
        outputs.append(dst)
        print(f'输出: {dst}')

    if incomplete:
        recs = []
        for gi, raw, room in incomplete:
            row = students_raw[gi]
            recs.append({
                '姓名': cell(row, col_idx['姓名']),
                '班级': cell(row, col_idx['班级']),
                '学号': cell(row, col_idx['学号']),
                '原填写内容': raw,
                '缺失项': '缺寝室楼号' if raw else '寝室信息为空',
            })
        default_name = f'信息不全人员-{stamp["y"]:04d}{stamp["m"]:02d}{stamp["d"]:02d}.docx'
        if args.name:
            stem, ext = os.path.splitext(args.name)
            name = f'{stem}-信息不全{ext or ".docx"}'
        else:
            name = default_name
        dst = os.path.join(args.out, name)
        write_incomplete_report(dst, recs)
        outputs.append(dst)
        print(f'\n=== 信息不全人员 {len(recs)} 人（缺寝室楼号，已单独成表）===')
        for rec in recs:
            print(f'  {rec["姓名"]}  {rec["班级"]}  {rec["学号"]}  原填写「{rec["原填写内容"]}」')
        print(f'输出: {dst}')

    print(f'\n落款日期: {sign_date}  |  楼号分组: {[norm_building(b) for b in sorted(groups)]}')
    return outputs


if __name__ == '__main__':
    main()
