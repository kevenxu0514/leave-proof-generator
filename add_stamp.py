# -*- coding: utf-8 -*-
"""
第二步：给 Word 初版加盖印章，并导出 PDF。

流程（对应「两步交付」）：
  第一步 generate.py 出 Word 初版 + 信息不全人员提示 → 人工复检
  第二步 本脚本：使用者提供章印图片 → 白底透明化 → 浮于文字上方压到落款日期处 → 导出 PDF

用法:
    python add_stamp.py --docx 分寝室请假证明-8舍-20261011.docx --seal 章.png

参数:
    --docx        第一步产出的 Word 初版
    --seal        章印图片（红章白底，png/jpg 均可）
    --pdf         输出 PDF 路径（默认与盖章后的 docx 同名）
    --docx-out    盖章后的 docx 路径（默认 <原名>_盖章.docx）
    --width-cm    章印宽度，默认 4.2cm，高度按原图比例
    --offset-x-cm 水平微调，负数左移，默认 0
    --offset-y-cm 垂直微调，负数上移，默认 0
    --first-page-only  只在第一页盖章（默认每一页落款处都盖）
"""
import argparse
import os
import re
import sys
import tempfile

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn, nsdecls
from docx.shared import Cm

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from generate import render_pdf          # 复用已有的 Word/WPS 渲染通道

EMU_PER_CM = 360000
DATE_RE = re.compile(r'\s*\d{4}年\d{1,2}月\d{1,2}日\s*$')


def make_white_transparent(seal_path, out_png, keep_red_only=True):
    """
    章印是红章白底。除了红色部分，其余一律设为透明，
    边缘按「红度」给半透明，避免锯齿。
    """
    try:
        from PIL import Image
    except ImportError:
        sys.exit('缺少依赖 Pillow，请先执行: pip install Pillow')

    im = Image.open(seal_path).convert('RGBA')
    w, h = im.size
    px = im.load()
    for y in range(h):
        for x in range(w):
            r, g, b, a = px[x, y]
            if keep_red_only:
                redness = min(r - g, r - b)
                if redness <= 10:
                    px[x, y] = (0, 0, 0, 0)
                    continue
                alpha = a if redness >= 50 else int(a * redness / 50)
                px[x, y] = (r, g, b, alpha)
            else:
                # 只做白底透明
                if r > 240 and g > 240 and b > 240:
                    px[x, y] = (r, g, b, 0)
    im.save(out_png)
    return out_png, im.size


def add_floating_image(paragraph, image_path, width_cm, offset_x_cm, offset_y_cm,
                       col_width_emu):
    """把图片插成「浮于文字上方」的浮动图形，锚定在指定段落"""
    run = paragraph.add_run()
    shape = run.add_picture(image_path, width=Cm(width_cm))
    inline = shape._inline

    extent = inline.find(qn('wp:extent'))
    cx, cy = int(extent.get('cx')), int(extent.get('cy'))
    docPr = inline.find(qn('wp:docPr'))
    graphic = inline.find(qn('a:graphic'))

    # 默认贴右：章的右缘距右边界约 0.6cm
    pos_h = col_width_emu - cx - int(0.6 * EMU_PER_CM) + int(offset_x_cm * EMU_PER_CM)
    # 默认上移约 0.62 个章高，压住两行落款单位名与日期
    pos_v = -int(cy * 0.62) + int(offset_y_cm * EMU_PER_CM)

    anchor = parse_xml(
        f'<wp:anchor {nsdecls("wp", "a", "r")} '
        f'distT="0" distB="0" distL="0" distR="0" simplePos="0" '
        f'relativeHeight="251658240" behindDoc="0" locked="0" '
        f'layoutInCell="1" allowOverlap="1">'
        f'<wp:simplePos x="0" y="0"/>'
        f'<wp:positionH relativeFrom="column"><wp:posOffset>{pos_h}</wp:posOffset></wp:positionH>'
        f'<wp:positionV relativeFrom="paragraph"><wp:posOffset>{pos_v}</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{cx}" cy="{cy}"/>'
        f'<wp:effectExtent l="0" t="0" r="0" b="0"/>'
        f'<wp:wrapNone/>'
        f'<wp:docPr id="{docPr.get("id")}" name="seal"/>'
        f'<wp:cNvGraphicFramePr/>'
        f'{graphic.xml}'
        f'</wp:anchor>')
    inline.getparent().replace(inline, anchor)
    return paragraph


def main():
    ap = argparse.ArgumentParser(description='给请假证明加盖印章并导出 PDF')
    ap.add_argument('--docx', required=True, help='第一步产出的 Word 初版')
    ap.add_argument('--seal', required=True, help='章印图片（红章白底）')
    ap.add_argument('--pdf', default=None, help='输出 PDF 路径')
    ap.add_argument('--docx-out', default=None, help='盖章后的 docx 路径')
    ap.add_argument('--width-cm', type=float, default=4.2, help='章印宽度 cm，默认 4.2')
    ap.add_argument('--offset-x-cm', type=float, default=0.0)
    ap.add_argument('--offset-y-cm', type=float, default=0.0)
    ap.add_argument('--first-page-only', action='store_true',
                    help='只在第一页盖章（默认每一页落款处都盖）')
    args = ap.parse_args()

    if not os.path.exists(args.docx):
        sys.exit(f'找不到 Word 初版: {args.docx}')
    if not os.path.exists(args.seal):
        sys.exit(f'找不到章印图片: {args.seal}')

    base = os.path.splitext(os.path.abspath(args.docx))[0]
    docx_out = args.docx_out or f'{base}_盖章.docx'
    pdf_out = args.pdf or f'{base}_盖章.pdf'

    # 1. 章印去白底
    tmp_png = os.path.join(tempfile.gettempdir(), '_seal_transparent.png')
    _, size = make_white_transparent(args.seal, tmp_png)
    print(f'章印处理: {args.seal} -> 白底透明化 ({size[0]}x{size[1]})')

    # 2. 插入浮动章
    doc = Document(args.docx)
    section = doc.sections[0]
    col_width = section.page_width - section.left_margin - section.right_margin
    dates = [p for p in doc.paragraphs if DATE_RE.match(p.text)]
    if not dates:
        sys.exit('文档里找不到落款日期段落，无法定位盖章位置')
    targets = dates[:1] if args.first_page_only else dates
    for p in targets:
        add_floating_image(p, tmp_png, args.width_cm,
                           args.offset_x_cm, args.offset_y_cm, col_width)
    doc.save(docx_out)
    print(f'盖章完成: {docx_out}（共 {len(targets)} 处落款）')

    # 3. 导出 PDF
    render_pdf(docx_out)
    produced = os.path.splitext(docx_out)[0] + '.pdf'
    if os.path.abspath(produced) != os.path.abspath(pdf_out):
        os.replace(produced, pdf_out)
    print(f'输出 PDF: {pdf_out}')


if __name__ == '__main__':
    main()
